import io

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None

try:
    from PIL import Image
    import pytesseract
except ImportError:
    Image = None
    pytesseract = None

def parse_pdf(file_bytes: bytes) -> str:
    """
    Extracts text from PDF, trying PyMuPDF, then pypdf, and OCR fallback if scanned.
    """
    text = ""
    # Try PyMuPDF first
    if fitz is not None:
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            for page in doc:
                page_text = page.get_text()
                if page_text:
                    text += page_text + "\n"

            # Check if scanned and OCR is available
            if len(text.strip()) < 30 and Image is not None and pytesseract is not None:
                ocr_text = ""
                for page in doc:
                    pix = page.get_pixmap(dpi=150)
                    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    try:
                        ocr_text += pytesseract.image_to_string(img) + "\n"
                    except Exception:
                        pass
                if len(ocr_text.strip()) > len(text.strip()):
                    text = ocr_text

            doc.close()
            if text.strip():
                return text.strip()
        except Exception as e:
            print(f"[Resume Parser] PyMuPDF error: {e}")

    # Fallback to pypdf
    if pypdf is not None:
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
            if text.strip():
                return text.strip()
        except Exception as e:
            print(f"[Resume Parser] pypdf error: {e}")

    return text.strip()

def parse_docx(file_bytes: bytes) -> str:
    """
    Extracts text from Microsoft Word .docx files.
    """
    if docx is None:
        print("[Resume Parser] python-docx is not installed.")
        return ""

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs).strip()
    except Exception as e:
        print(f"[Resume Parser] DOCX parsing error: {e}")
        return ""

def parse_image(file_bytes: bytes) -> str:
    """
    Extracts text from images using OCR.
    """
    if Image is None or pytesseract is None:
        print("[Resume Parser] PIL or pytesseract not available.")
        return ""

    try:
        img = Image.open(io.BytesIO(file_bytes))
        return pytesseract.image_to_string(img).strip()
    except Exception as e:
        print(f"[Resume Parser] Image OCR error: {e}")
        return ""

def parse_resume_file(file_storage) -> str:
    """
    Main dispatcher for extracting text from uploaded resume file.
    Accepts a Werkzeug FileStorage object.
    """
    filename = file_storage.filename or ""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    file_bytes = file_storage.read()

    # Reset stream position
    file_storage.seek(0)

    if ext == "pdf":
        return parse_pdf(file_bytes)
    elif ext in {"docx", "doc"}:
        return parse_docx(file_bytes)
    elif ext in {"png", "jpg", "jpeg"}:
        return parse_image(file_bytes)
    elif ext == "txt":
        try:
            return file_bytes.decode("utf-8", errors="ignore").strip()
        except Exception:
            return ""
    else:
        # Fallback trial
        text = parse_pdf(file_bytes)
        if not text:
            text = parse_docx(file_bytes)
        return text
